"""Contrôles structurels des livrables Word (sans dépendance à un moteur de rendu)."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
EXPECTED = {
    'Cours_Biochimie_L2_Structure_des_glucides.docx': {'pages': 23, 'images': 17},
    'Corrige_detaille_TD1_Biochimie_L2.docx': {'pages': 19, 'images': 12},
}


def check(path, expected):
    with ZipFile(path) as z:
        assert z.testzip() is None, 'Archive DOCX corrompue'
        root = etree.fromstring(z.read('word/document.xml'))
        breaks = root.xpath('//w:pPr/w:pageBreakBefore', namespaces=NS)
        anchors = set(root.xpath('//w:bookmarkStart/@w:name', namespaces=NS))
        links = set(root.xpath('//w:hyperlink/@w:anchor', namespaces=NS))
        assert not links - anchors, f'Liens internes cassés : {links - anchors}'
        assert len(breaks) + 1 == expected['pages'], 'Nombre de pages prévues incohérent'
        assert 'word/styles.xml' in z.namelist()
    d = Document(path)
    assert len(d.inline_shapes) == expected['images'], 'Nombre de figures intégrées incohérent'
    text = '\n'.join(p.text for p in d.paragraphs)
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                text += '\n' + cell.text
    for marker in ['D-glucose', 'D-fructose', 'saccharose', 'lactose', 'maltose', 'C₁₂H₂₂O₁₁']:
        assert marker.lower() in text.lower(), f'Contenu attendu absent : {marker}'
    for bad in ['TODO', 'Lorem ipsum', 'Evaluation Only', 'Aspose']:
        assert bad not in text, f'Marqueur indésirable : {bad}'
    print(f'OK — {path.name} : {expected["pages"]} pages prévues, {len(d.inline_shapes)} figures, {len(d.tables)} tableaux.')


def main():
    for name, expected in EXPECTED.items():
        check(ROOT / 'documents' / name, expected)


if __name__ == '__main__':
    main()
